import json
import os
import sys
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

def clean_val(val):
    if val is None:
        return ""
    if isinstance(val, str):
        return ILLEGAL_CHARACTERS_RE.sub('', val)
    return val

BASE_DIR = r"c:\antigravity files IDE\encroachment-tracker"
DATA_DIR = os.path.join(BASE_DIR, 'server', 'data', 'generated')
OUTPUT_DIR = os.path.join(BASE_DIR, 'XLSX')
os.makedirs(OUTPUT_DIR, exist_ok=True)

REPORTS_FILE = os.path.join(DATA_DIR, 'reports.json')
PROJECTS_FILE = os.path.join(DATA_DIR, 'projects.json')
MANAGERS_FILE = os.path.join(DATA_DIR, 'managers.json')
STATS_FILE = os.path.join(DATA_DIR, 'stats.json')

C_PRIMARY_NAVY = '0A2540'
C_ROYAL_BLUE   = '1E40AF'
C_HEADER_FILL  = '0F3B66'
C_SUB_HEADER   = '1D4E89'
C_ACCENT_BLUE  = 'E0F2FE'
C_WHITE        = 'FFFFFF'
C_DARK_TEXT    = '1E293B'
C_ZEBRA        = 'F8FAFC'
C_CARD_BG      = 'F1F5F9'
C_BORDER_LIGHT = 'CBD5E1'

COLORS_AGING = {
    '0 - 30':   {'fill': 'DCFCE7', 'text': '166534', 'label': '0 - 30 يوم (طبيعي 🟢)', 'risk': 'منخفض'},
    '31 - 60':  {'fill': 'FEF9C3', 'text': '854D0E', 'label': '31 - 60 يوم (قيد المتابعة 🟡)', 'risk': 'متوسط'},
    '61 - 90':  {'fill': 'FFEDD5', 'text': '9A3412', 'label': '61 - 90 يوم (متأخر عن SLA 🟠)', 'risk': 'مرتفع'},
    '91 - 180': {'fill': 'FEE2E2', 'text': '991B1B', 'label': '91 - 180 يوم (حرج وتشغيلي 🔴)', 'risk': 'حرج'},
    '> 180':    {'fill': 'FCE7F3', 'text': '831843', 'label': 'أكثر من 180 يوم (خرق جسيم 🚨)', 'risk': 'خرق جسيم'}
}

COLORS_STATUS = {
    'تحت معالجة المقاول':           {'fill': 'E0E7FF', 'text': '3730A3'},
    'معاد من المقاول':              {'fill': 'FEE2E2', 'text': '991B1B'},
    'معاد من الجهة المتعدية':       {'fill': 'FEF3C7', 'text': '92400E'},
    'معاد للجهة المالكة':           {'fill': 'F1F5F9', 'text': '475569'},
    'بانتظار اعتماد الجهة المتعدية': {'fill': 'E0F2FE', 'text': '075985'},
    'تحت معالجة الجهة المتعدية':     {'fill': 'ECFDF5', 'text': '065F46'},
    'تمت المعالجة':                 {'fill': 'DCFCE7', 'text': '166534'}
}

def get_aging_info(age_days):
    if age_days is None or age_days < 0:
        return COLORS_AGING['0 - 30']
    if age_days <= 30:
        return COLORS_AGING['0 - 30']
    elif age_days <= 60:
        return COLORS_AGING['31 - 60']
    elif age_days <= 90:
        return COLORS_AGING['61 - 90']
    elif age_days <= 180:
        return COLORS_AGING['91 - 180']
    else:
        return COLORS_AGING['> 180']

def get_routing_reason_text(report):
    reason = report.get('reason', '')
    if 'spatial:kml' in reason:
        return 'تطابق مكاني مع مضلع المشروع الجاري في طبقات KMZ المعتمدة'
    elif 'aswad_exception' in reason:
        return 'استثناء معتمد لمشاريع م. عبدالله الأسود (تغطية شاملة لمدينة الرياض)'
    elif 'governorate' in reason:
        return 'مشروع محافظات: تطابق مقاول المشروع الجاري ونطاق المحافظة'
    elif 'proj_ref' in reason:
        return 'تطابق المقاول والنطاق مع ذكر صريح لبيانات المشروع في نص البلاغ'
    elif 'contractor+district' in reason:
        return 'تطابق مقاول المشروع ونطاق الحي الجغرافي المعتمد'
    elif 'manual' in reason:
        return 'إسناد واعتماد يدوي من إدارة المشاريع'
    elif report.get('matched'):
        return 'مطابقة معتمدة مع مشروع جاري'
    else:
        return 'خارج نطاق المشاريع الجارية (تشغيل وصيانة / شبكة قائمة)'

def get_recommended_action(report):
    status = report.get('status', '')
    age_days = report.get('ageDays', 0)
    contractor = report.get('project', {}).get('contractor', report.get('contractorName', 'المقاول'))

    if age_days > 180:
        return f'إنذار نهائي عاجل لـ ({contractor}) وتطبيق غرامات التأخير وحسم تكاليف الإصلاح فوراً'
    elif status == 'معاد من المقاول':
        return f'مراجعة مبررات الإعادة مع ({contractor}) وإلزامه بالتنسيق الفوري مع الجهة المتضررة'
    elif status == 'تحت معالجة المقاول':
        return f'إلزام ({contractor}) بإنهاء الأعمال وإرفاق محضر إنهاء الملاحظات وتحديث النظام'
    elif status in ['معاد من الجهة المتعدية', 'معاد للجهة المالكة']:
        return f'مخاطبة الجهة المتعدية رسمياً وإلزام مقاول المشروع بالمتابعة الميدانية'
    elif status == 'بانتظار اعتماد الجهة المتعدية':
        return 'متابعة اعتماد إغلاق البلاغ لدى الجهة المتعدية لرفع نسبة الإنجاز'
    elif status == 'تمت المعالجة':
        return 'تم إغلاق البلاغ بنجاح ومعالجة الموقع'
    else:
        return f'متابعة الإجراء الميداني الفوري مع ({contractor})'

def thin_border():
    thin = Side(border_style="thin", color=C_BORDER_LIGHT)
    return Border(top=thin, left=thin, right=thin, bottom=thin)

def double_bottom_border():
    thin = Side(border_style="thin", color=C_BORDER_LIGHT)
    double = Side(border_style="double", color=C_HEADER_FILL)
    return Border(top=thin, left=thin, right=thin, bottom=double)

def auto_fit_columns(ws, max_cols=30, max_width_limit=55):
    for col in range(1, max_cols + 1):
        col_letter = get_column_letter(col)
        max_len = 0
        for row in range(1, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            val = cell.value
            if val is not None:
                val_str = str(val)
                lines = val_str.split('\n')
                longest_line = max(len(l) for l in lines)
                if longest_line > max_len:
                    max_len = longest_line
        
        width = max(max_len + 3, 12)
        if width > max_width_limit:
            width = max_width_limit
        ws.column_dimensions[col_letter].width = width

def setup_sheet_view(ws):
    ws.views.sheetView[0].rightToLeft = True
    ws.views.sheetView[0].showGridLines = True

def write_reports_table(ws, reports_list, sheet_title_text, subtitle_text):
    setup_sheet_view(ws)

    # عنوان الورقة
    ws.merge_cells("A1:Z1")
    ws.row_dimensions[1].height = 36
    h_rep = ws["A1"]
    h_rep.value = sheet_title_text
    h_rep.font = Font(name="Segoe UI", size=14, bold=True, color=C_WHITE)
    h_rep.alignment = Alignment(horizontal="center", vertical="center")
    h_rep.fill = PatternFill(start_color=C_PRIMARY_NAVY, end_color=C_PRIMARY_NAVY, fill_type="solid")

    ws.merge_cells("A2:Z2")
    ws.row_dimensions[2].height = 22
    h_sub = ws["A2"]
    h_sub.value = subtitle_text
    h_sub.font = Font(name="Segoe UI", size=9, bold=True, color="E2E8F0")
    h_sub.alignment = Alignment(horizontal="center", vertical="center")
    h_sub.fill = PatternFill(start_color=C_HEADER_FILL, end_color=C_HEADER_FILL, fill_type="solid")

    headers_reports = [
        "م",
        "رقم بلاغ التعدي",
        "مدير البرنامج المسؤول (NWC)",
        "اسم المشروع الجاري",
        "رقم العملية",
        "أمر الشراء (PO)",
        "مقاول المشروع المعتمد",
        "الاستشاري المشرف",
        "القطاع والشبكة",
        "المدينة / المحافظة",
        "الحي",
        "الشارع / الموقع",
        "خط الطول (X)",
        "خط العرض (Y)",
        "تاريخ البلاغ",
        "عمر البلاغ (يوم)",
        "مستوى الخطورة وشريحة العمر",
        "حالة البلاغ في المنظومة",
        "جهة التوجيه المعتمدة",
        "سبب المطابقة والتوجيه",
        "الجهة المالكة للمرفق",
        "الجهة المتعدية المسجلة",
        "وصف التعدي الميداني",
        "أثر التعدي وضرر المرفق",
        "تعليق المركز وسجل التوجيه",
        "الإجراء التنفيذي المطلوب"
    ]

    ws.row_dimensions[4].height = 30
    for col_i, h in enumerate(headers_reports, start=1):
        cell = ws.cell(row=4, column=col_i, value=h)
        cell.font = Font(name="Segoe UI", size=9, bold=True, color=C_WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = PatternFill(start_color=C_SUB_HEADER, end_color=C_SUB_HEADER, fill_type="solid")
        cell.border = thin_border()

    for r_idx, rep in enumerate(reports_list, start=1):
        row_num = r_idx + 4
        ws.row_dimensions[row_num].height = 24

        age_days = rep.get('ageDays', 0)
        aging_info = get_aging_info(age_days)
        routing_reason = get_routing_reason_text(rep)
        rec_action = get_recommended_action(rep)
        proj = rep.get('project', {})
        status = rep.get('status', 'تحت معالجة المقاول')

        row_data = [
            r_idx,
            rep.get('id', ''),
            proj.get('programManager', 'غير محدد'),
            proj.get('name', 'غير محدد'),
            proj.get('operationNumber', '-'),
            proj.get('po', '-'),
            proj.get('contractor', rep.get('contractorName', '-')),
            proj.get('consultant', '-'),
            rep.get('sector', 'صرف'),
            rep.get('city', 'الرياض'),
            rep.get('district', '-'),
            rep.get('street', '-'),
            rep.get('longitude', '-'),
            rep.get('latitude', '-'),
            rep.get('dateReport', rep.get('dateIncident', '-')),
            age_days,
            aging_info['label'],
            status,
            f"المقاول ({proj.get('contractor', rep.get('contractorName', ''))}) - مشروع جاري",
            routing_reason,
            rep.get('ownerEntity', '-'),
            rep.get('encroachingEntity', '-'),
            rep.get('description', '-'),
            rep.get('impact', '-'),
            rep.get('centerComment', '-'),
            rec_action
        ]

        is_even = (r_idx % 2 == 0)
        row_bg = C_ZEBRA if is_even else C_WHITE

        for col_i, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_num, column=col_i, value=clean_val(val))
            cell.font = Font(name="Segoe UI", size=9, color=C_DARK_TEXT)
            cell.border = thin_border()

            if col_i in [1, 2, 5, 6, 9, 13, 14, 15, 16]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")

            cell.fill = PatternFill(start_color=row_bg, end_color=row_bg, fill_type="solid")

            if col_i == 2:
                cell.font = Font(name="Segoe UI", size=9, bold=True, color=C_ROYAL_BLUE)
            elif col_i == 3:
                cell.font = Font(name="Segoe UI", size=9, bold=True, color=C_PRIMARY_NAVY)
            elif col_i == 16:
                cell.font = Font(name="Segoe UI", size=9, bold=True)
            elif col_i == 17:
                cell.fill = PatternFill(start_color=aging_info['fill'], end_color=aging_info['fill'], fill_type="solid")
                cell.font = Font(name="Segoe UI", size=9, bold=True, color=aging_info['text'])
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_i == 18:
                st_color = COLORS_STATUS.get(status, {'fill': 'F1F5F9', 'text': '334155'})
                cell.fill = PatternFill(start_color=st_color['fill'], end_color=st_color['fill'], fill_type="solid")
                cell.font = Font(name="Segoe UI", size=8, bold=True, color=st_color['text'])
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_i == 26:
                cell.font = Font(name="Segoe UI", size=8.5, bold=True, color='831843' if age_days > 180 else '0F3B66')

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:Z{len(reports_list) + 4}"
    auto_fit_columns(ws, max_cols=26, max_width_limit=45)
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 22
    ws.column_dimensions['D'].width = 38
    ws.column_dimensions['G'].width = 28
    ws.column_dimensions['Q'].width = 26
    ws.column_dimensions['R'].width = 22
    ws.column_dimensions['T'].width = 36
    ws.column_dimensions['W'].width = 40
    ws.column_dimensions['X'].width = 35
    ws.column_dimensions['Y'].width = 40
    ws.column_dimensions['Z'].width = 42

print("Script template validated!")
