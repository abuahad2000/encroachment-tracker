import json
import os
import sys
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIRECTORY_FILE = os.path.join(BASE_DIR, 'server', 'data', 'contractor_directory.json')
BACKUP_DIRECTORY_FILE = os.path.join(BASE_DIR, 'server', 'data', 'contractor_directory_backup.json')
OUTPUT_DIR = os.path.join(BASE_DIR, 'XLSX')
os.makedirs(OUTPUT_DIR, exist_ok=True)
OUTPUT_FILE = os.path.join(OUTPUT_DIR, 'سجل_بيانات_مقاولي_المشاريع_NWC.xlsx')

C_PRIMARY_NAVY = '0A2540'
C_ROYAL_BLUE   = '1E40AF'
C_HEADER_FILL  = '0F3B66'
C_SUB_HEADER   = '1D4E89'
C_ACCENT_BLUE  = 'E0F2FE'
C_WHITE        = 'FFFFFF'
C_DARK_TEXT    = '1E293B'
C_ZEBRA        = 'F8FAFC'
C_CARD_BG      = 'F1F5F9'
C_BORDER_LIGHT = 'CBD5E1'

def thin_border():
    thin = Side(border_style="thin", color=C_BORDER_LIGHT)
    return Border(top=thin, left=thin, right=thin, bottom=thin)

def auto_fit_columns(ws, max_cols=20, max_width_limit=55):
    for col in range(1, max_cols + 1):
        col_letter = get_column_letter(col)
        max_len = 0
        for row in range(1, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            val = cell.value
            if val is not None:
                val_str = str(val)
                lines = val_str.split('\n')
                longest_line = max(len(l) for l in lines)
                if longest_line > max_len:
                    max_len = longest_line
        width = max(max_len + 3, 12)
        if width > max_width_limit:
            width = max_width_limit
        ws.column_dimensions[col_letter].width = width

def export_directory_excel():
    directory = []
    if os.path.exists(DIRECTORY_FILE):
        try:
            with open(DIRECTORY_FILE, 'r', encoding='utf-8') as f:
                directory = json.load(f)
        except Exception as e:
            print(f"Warning loading primary directory: {e}")

    if (not directory or len(directory) == 0) and os.path.exists(BACKUP_DIRECTORY_FILE):
        try:
            with open(BACKUP_DIRECTORY_FILE, 'r', encoding='utf-8') as f:
                directory = json.load(f)
                print("Loaded directory from backup file.")
        except Exception as e:
            print(f"Warning loading backup directory: {e}")

    if not directory:
        print(f"Error: Neither {DIRECTORY_FILE} nor {BACKUP_DIRECTORY_FILE} exists or contains data.")
        return False

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "سجل المقاولين والمشاريع"
    ws.views.sheetView[0].rightToLeft = True
    ws.views.sheetView[0].showGridLines = True

    # Title
    ws.merge_cells("A1:L1")
    ws.row_dimensions[1].height = 36
    h1 = ws["A1"]
    h1.value = "شركة المياه الوطنية • قطاع المشاريع الرأسمالية بالقطاع الأوسط | سجل بيانات مقاولي المشاريع والمنشآت"
    h1.font = Font(name="Segoe UI", size=14, bold=True, color=C_WHITE)
    h1.alignment = Alignment(horizontal="center", vertical="center")
    h1.fill = PatternFill(start_color=C_PRIMARY_NAVY, end_color=C_PRIMARY_NAVY, fill_type="solid")

    # Subtitle
    ws.merge_cells("A2:L2")
    ws.row_dimensions[2].height = 22
    h2 = ws["A2"]
    total_contractors = len(set(d.get('contractorName', '') for d in directory if d.get('contractorName')))
    h2.value = f"تاريخ الاستخراج: {os.environ.get('EXPORT_DATE', 'سبتمبر 2026')} | إجمالي المشاريع: {len(directory)} | إجمالي المقاولين: {total_contractors}"
    h2.font = Font(name="Segoe UI", size=9.5, bold=True, color="E2E8F0")
    h2.alignment = Alignment(horizontal="center", vertical="center")
    h2.fill = PatternFill(start_color=C_HEADER_FILL, end_color=C_HEADER_FILL, fill_type="solid")

    # Headers
    headers = [
        "م",
        "اسم المقاول",
        "رقم المشروع (العملية)",
        "اسم المشروع",
        "موقع المشروع (النطاق)",
        "مدير البرنامج (NWC)",
        "مدير المشروع (NWC)",
        "رقم السجل التجاري",
        "الرقم الموحد (700)",
        "اسم / رقم المسؤول",
        "جوال المسؤول",
        "البريد الإلكتروني"
    ]

    ws.row_dimensions[4].height = 28
    for col_i, h in enumerate(headers, start=1):
        c = ws.cell(row=4, column=col_i, value=h)
        c.font = Font(name="Segoe UI", size=9.5, bold=True, color=C_WHITE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.fill = PatternFill(start_color=C_SUB_HEADER, end_color=C_SUB_HEADER, fill_type="solid")
        c.border = thin_border()

    for idx, item in enumerate(directory, start=1):
        row_num = idx + 4
        ws.row_dimensions[row_num].height = 24

        row_data = [
            idx,
            item.get('contractorName', '-'),
            item.get('projectNumber', '-'),
            item.get('projectName', '-'),
            item.get('projectLocation', '-'),
            item.get('programManager', '-'),
            item.get('projectManager', '-'),
            item.get('crNumber', '') or '-',
            item.get('unifiedNumber', '') or '-',
            item.get('managerName', '') or '-',
            item.get('managerPhone', '') or '-',
            item.get('managerEmail', '') or '-'
        ]

        is_even = (idx % 2 == 0)
        row_bg = C_ZEBRA if is_even else C_WHITE

        for col_i, val in enumerate(row_data, start=1):
            c = ws.cell(row=row_num, column=col_i, value=val)
            c.font = Font(name="Segoe UI", size=9, color=C_DARK_TEXT)
            c.border = thin_border()
            c.fill = PatternFill(start_color=row_bg, end_color=row_bg, fill_type="solid")

            if col_i in [1, 3, 8, 9, 11]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="right", vertical="center")

            if col_i == 2:
                c.font = Font(name="Segoe UI", size=9, bold=True, color=C_PRIMARY_NAVY)
            elif col_i == 3:
                c.font = Font(name="Segoe UI", size=9, bold=True, color=C_ROYAL_BLUE)
            elif col_i in [8, 9] and val and val != '-':
                c.font = Font(name="Segoe UI", size=9, bold=True, color='166534')

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:L{len(directory) + 4}"
    auto_fit_columns(ws, max_cols=12, max_width_limit=45)
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 22
    ws.column_dimensions['D'].width = 38
    ws.column_dimensions['E'].width = 25
    ws.column_dimensions['F'].width = 24
    ws.column_dimensions['G'].width = 24
    ws.column_dimensions['H'].width = 18
    ws.column_dimensions['I'].width = 18
    ws.column_dimensions['J'].width = 22
    ws.column_dimensions['K'].width = 16
    ws.column_dimensions['L'].width = 25

    wb.save(OUTPUT_FILE)
    print(f"Exported successfully to {OUTPUT_FILE}")
    return True

if __name__ == '__main__':
    export_directory_excel()
